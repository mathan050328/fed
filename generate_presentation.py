"""
Automated generation of a professional PowerPoint presentation (.pptx)
for the Multi-Disease Federated Learning System (Heart Disease, Diabetes, CKD).
"""

from pathlib import Path
import json
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = Path(__file__).resolve().parent
ARTIFACTS_DIR = BASE_DIR / "artifacts"
FIGURES_DIR = ARTIFACTS_DIR / "figures"
METRICS_DIR = ARTIFACTS_DIR / "metrics"
PPTX_PATH = BASE_DIR / "Federated_Disease_Prediction_Presentation.pptx"

# Theme Colors
COLOR_NAVY = RGBColor(26, 54, 93)       # #1A365D
COLOR_TEAL = RGBColor(14, 116, 144)     # #0E7490
COLOR_DARK_GRAY = RGBColor(51, 65, 85)  # #334155
COLOR_LIGHT_BG = RGBColor(248, 250, 252)# #F8FAFC
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_CARD_BORDER = RGBColor(226, 232, 240)


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="MULTI-DISEASE FEDERATED LEARNING PLATFORM"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_c = cat_box.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = COLOR_TEAL

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY

    def add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # SLIDE 1: Title
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_NAVY
    bg1.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(2.2))
    tf = t_box.text_frame
    p0 = tf.paragraphs[0]
    p0.text = "ACADEMIC FINAL-YEAR CAPSTONE PROJECT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(56, 189, 248)

    p1 = tf.add_paragraph()
    p1.text = "Multi-Disease Federated Learning\nPrediction Platform"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_before = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "Collaborative Diagnosis for Heart Disease, Diabetes, and CKD across Hospitals A, B, and C"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_before = Pt(14)

    # SLIDE 2: Problem Statement
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Statement: Siloed Healthcare Intelligence")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1), COLOR_LIGHT_BG)
    b1 = s2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.7)).text_frame
    b1.word_wrap = True
    p = b1.paragraphs[0]
    p.text = "Centralized Healthcare Bottlenecks"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_NAVY
    bullets = [
        "Multiple Clinical Silos: Hospitals treat diverse diseases but cannot pool records due to HIPAA/GDPR.",
        "Single-Disease Narrow Focus: Existing AI systems often only predict one disease, requiring redundant infrastructure.",
        "Central Warehouse Security Threat: Merging cardiology, endocrinology, and nephrology records creates a massive attack vector.",
    ]
    for b in bullets:
        p_b = b1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13); p_b.space_before = Pt(12)

    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.1), RGBColor(240, 253, 250), COLOR_TEAL)
    b2 = s2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.7)).text_frame
    b2.word_wrap = True
    p = b2.paragraphs[0]
    p.text = "Multi-Disease Federated Architecture"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = COLOR_TEAL
    bullets2 = [
        "Local Multi-Disease DBs: Hospital A, B, and C keep Heart, Diabetes, and CKD data in strictly isolated local databases.",
        "Independent Local Optimization: Each hospital trains local models for each disease domain separately.",
        "Model Updates Only: Transmits abstract parameter weights (coef_, intercept_) to central Flower coordinator.",
        "Unified Healthcare UI: Single clinical dashboard supporting disease selection, prediction, and feature explanations.",
    ]
    for b in bullets2:
        p_b = b2.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(13); p_b.space_before = Pt(12)

    # SLIDE 3: System Architecture (Matching User Diagram)
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "System Topology: Federated Server & Hospital Nodes")

    # Top: Federated Server
    add_card(s3, Inches(3.6), Inches(1.5), Inches(6.1), Inches(1.1), RGBColor(224, 242, 254), COLOR_TEAL)
    tf_top = s3.shapes.add_textbox(Inches(3.8), Inches(1.55), Inches(5.7), Inches(0.9)).text_frame
    p = tf_top.paragraphs[0]; p.text = "FEDERATED SERVER"; p.font.size = Pt(14); p.font.bold = True; p.alignment = PP_ALIGN.CENTER; p.font.color.rgb = COLOR_NAVY
    p2 = tf_top.add_paragraph(); p2.text = "Flower Coordinator | FedAvg | Model Registry | Training Monitor"; p2.font.size = Pt(11); p2.alignment = PP_ALIGN.CENTER; p2.font.color.rgb = COLOR_DARK_GRAY

    # Middle: 3 Hospitals
    h_cards = [
        ("HOSPITAL A", "Local DB: Heart, Diabetes, CKD\nLocal Training", Inches(0.8), RGBColor(220, 252, 231)),
        ("HOSPITAL B", "Local DB: Heart, Diabetes, CKD\nLocal Training", Inches(5.0), RGBColor(254, 243, 199)),
        ("HOSPITAL C", "Local DB: Heart, Diabetes, CKD\nLocal Training", Inches(9.2), RGBColor(252, 231, 243)),
    ]
    for h_name, h_desc, left_p, bg_c in h_cards:
        add_card(s3, left_p, Inches(2.9), Inches(3.3), Inches(1.4), bg_c)
        tf_h = s3.shapes.add_textbox(left_p + Inches(0.1), Inches(3.0), Inches(3.1), Inches(1.2)).text_frame
        p = tf_h.paragraphs[0]; p.text = h_name; p.font.size = Pt(13); p.font.bold = True; p.alignment = PP_ALIGN.CENTER; p.font.color.rgb = COLOR_NAVY
        p2 = tf_h.add_paragraph(); p2.text = h_desc; p2.font.size = Pt(10.5); p2.alignment = PP_ALIGN.CENTER; p2.font.color.rgb = COLOR_DARK_GRAY

    # Aggregation Label
    b_agg = s3.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.5)).text_frame
    p_agg = b_agg.paragraphs[0]; p_agg.text = "⬇ [ MODEL UPDATES ONLY ] ⬇  [ FEDERATED AGGREGATION via FedAvg ]  ⬇"; p_agg.font.size = Pt(11); p_agg.font.bold = True; p_agg.alignment = PP_ALIGN.CENTER; p_agg.font.color.rgb = COLOR_TEAL

    # Bottom: Global Models & UI
    add_card(s3, Inches(2.5), Inches(5.1), Inches(8.3), Inches(1.0), RGBColor(243, 232, 255))
    tf_gm = s3.shapes.add_textbox(Inches(2.7), Inches(5.15), Inches(7.9), Inches(0.8)).text_frame
    p = tf_gm.paragraphs[0]; p.text = "GLOBAL MODELS REGISTRY"; p.font.size = Pt(13); p.font.bold = True; p.alignment = PP_ALIGN.CENTER; p.font.color.rgb = COLOR_NAVY
    p2 = tf_gm.add_paragraph(); p2.text = "Heart Disease Model (85.25%) | Diabetes Model (71.43%) | CKD Model (95.00%)"; p2.font.size = Pt(11); p2.alignment = PP_ALIGN.CENTER

    add_card(s3, Inches(2.5), Inches(6.25), Inches(8.3), Inches(0.8), RGBColor(254, 249, 195))
    tf_ui = s3.shapes.add_textbox(Inches(2.7), Inches(6.3), Inches(7.9), Inches(0.7)).text_frame
    p = tf_ui.paragraphs[0]; p.text = "AI HEALTHCARE UI"; p.font.size = Pt(12); p.font.bold = True; p.alignment = PP_ALIGN.CENTER; p.font.color.rgb = COLOR_NAVY
    p2 = tf_ui.add_paragraph(); p2.text = "Disease Selection | Patient Input | Prediction | Risk Probability | Explanation | Model Metrics"; p2.font.size = Pt(10.5); p2.alignment = PP_ALIGN.CENTER

    # SLIDE 4: Multi-Disease Benchmarks
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Multi-Disease Clinical Benchmarks")
    ds_cards = [
        ("🫀 Heart Disease", "303 Patient Records\n13 Clinical Features\nBinary Risk (0/1)\nKey: ST Depression, Chest Pain, Cholesterol", Inches(0.8)),
        ("🩸 Diabetes", "768 Patient Records\n8 Clinical Features\nBinary Risk (0/1)\nKey: Glucose, Insulin, BMI, Pedigree", Inches(5.0)),
        ("🩺 Chronic Kidney Disease", "400 Patient Records\n20 Clinical Features\nBinary Risk (0/1)\nKey: Serum Creatinine, Albumin, Urea, Hemoglobin", Inches(9.2)),
    ]
    for d_title, d_body, l_pos in ds_cards:
        add_card(s4, l_pos, Inches(1.6), Inches(3.3), Inches(4.8), COLOR_LIGHT_BG)
        tf_d = s4.shapes.add_textbox(l_pos + Inches(0.2), Inches(1.8), Inches(2.9), Inches(4.3)).text_frame
        p = tf_d.paragraphs[0]; p.text = d_title; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = COLOR_NAVY
        p2 = tf_d.add_paragraph(); p2.text = d_body; p2.font.size = Pt(12.5); p2.space_before = Pt(14); p2.font.color.rgb = COLOR_DARK_GRAY

    # SLIDE 5: Multi-Disease Experimental Results
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Multi-Disease Experimental Results: Centralized vs Federated")

    table_shape = s5.shapes.add_table(5, 6, Inches(0.8), Inches(1.6), Inches(11.7), Inches(3.6))
    t = table_shape.table
    col_w = [Inches(2.5), Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.8), Inches(2.0)]
    for idx, w in enumerate(col_w):
        t.columns[idx].width = w

    headers = ["Disease Domain", "Centralized Acc", "Federated Acc", "Centralized AUC", "Federated AUC", "Privacy Status"]
    for i, h in enumerate(headers):
        cell = t.cell(0, i); cell.text = h; cell.fill.solid(); cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(11); p.font.color.rgb = COLOR_WHITE

    res_data = [
        ("Heart Disease", "86.89%", "85.25%", "0.9102", "0.8950", "0 Records Shared"),
        ("Diabetes", "71.43%", "71.43%", "0.8230", "0.8241", "0 Records Shared"),
        ("Chronic Kidney Disease", "97.50%", "95.00%", "0.9901", "1.0000", "0 Records Shared"),
        ("AVERAGE OVERALL", "85.27%", "83.89%", "0.9078", "0.9064", "100% Data Decentralized"),
    ]
    for r_idx, row in enumerate(res_data, start=1):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx); cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(240, 253, 250) if r_idx == 4 else (COLOR_LIGHT_BG if r_idx % 2 == 0 else COLOR_WHITE)
            p = cell.text_frame.paragraphs[0]; p.font.size = Pt(11); p.font.color.rgb = COLOR_DARK_GRAY
            if r_idx == 4: p.font.bold = True

    add_card(s5, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.2), RGBColor(240, 253, 250), COLOR_TEAL)
    tf_call = s5.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.0)).text_frame
    p = tf_call.paragraphs[0]; p.text = "Key Multi-Disease Finding:"; p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = COLOR_TEAL
    p2 = tf_call.add_paragraph()
    p2.text = "Across all three diverse clinical domains (Cardiology, Endocrinology, and Nephrology), Federated Learning achieved comparable diagnostic discrimination (90.64% average ROC-AUC vs 90.78% centralized) with ZERO raw patient records leaving Hospital A, B, or C."
    p2.font.size = Pt(11); p2.font.color.rgb = COLOR_DARK_GRAY; p2.space_before = Pt(4)

    # SLIDE 6: UI & Explainability
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "AI Healthcare UI & Feature Explainability")
    ui_cards = [
        ("Disease Selection", "Seamless switching between Heart Disease, Diabetes, and Chronic Kidney Disease.", Inches(0.8)),
        ("Dynamic Patient Input", "Interactive forms adapted to the specific clinical parameters of the selected domain.", Inches(3.8)),
        ("Risk Probability Output", "Real-time probability gauge with clinical risk classification alerts.", Inches(6.8)),
        ("Model Explanation", "Biomarker contribution analysis highlighting top clinical risk drivers (odds ratios).", Inches(9.8)),
    ]
    for title, desc, left_pos in ui_cards:
        add_card(s6, left_pos, Inches(1.6), Inches(2.7), Inches(4.8), COLOR_LIGHT_BG)
        box = s6.shapes.add_textbox(left_pos + Inches(0.15), Inches(1.8), Inches(2.4), Inches(4.3)).text_frame
        box.word_wrap = True
        p = box.paragraphs[0]; p.text = title; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = COLOR_NAVY
        p2 = box.add_paragraph(); p2.text = desc; p2.font.size = Pt(11.5); p2.space_before = Pt(10); p2.font.color.rgb = COLOR_DARK_GRAY

    # SLIDE 7: Summary & Conclusion
    s7 = prs.slides.add_slide(blank_layout)
    bg7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg7.fill.solid(); bg7.fill.fore_color.rgb = COLOR_NAVY; bg7.line.fill.background()

    tf_c = s7.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(11.0), Inches(5.0)).text_frame
    p0 = tf_c.paragraphs[0]; p0.text = "SUMMARY & CONCLUSION"; p0.font.size = Pt(13); p0.font.bold = True; p0.font.color.rgb = RGBColor(56, 189, 248)
    p1 = tf_c.add_paragraph(); p1.text = "Multi-Disease Federated Learning Validated"; p1.font.size = Pt(30); p1.font.bold = True; p1.font.color.rgb = COLOR_WHITE; p1.space_before = Pt(8)

    bullets = [
        "3 Medical Conditions: Successfully deployed and registered global models for Heart Disease, Diabetes, and CKD.",
        "3 Hospital Partitions: Isolated local databases for Hospital A, B, and C with zero cross-institutional data sharing.",
        "Flower + FedAvg: Robust orchestration across 5 communication rounds achieving high ROC-AUC (89.5% Heart, 82.4% Diabetes, 100% CKD).",
        "Clinical AI UI: End-to-end web dashboard featuring disease switching, risk prediction, explanation, and model registry.",
    ]
    for b in bullets:
        p_b = tf_c.add_paragraph(); p_b.text = "✔  " + b; p_b.font.size = Pt(14); p_b.font.color.rgb = RGBColor(226, 232, 240); p_b.space_before = Pt(12)

    p_qa = tf_c.add_paragraph(); p_qa.text = "Thank You! Questions & Discussion Welcome."; p_qa.font.size = Pt(20); p_qa.font.bold = True; p_qa.font.color.rgb = RGBColor(56, 189, 248); p_qa.space_before = Pt(24)

    prs.save(PPTX_PATH)
    print(f"Updated multi-disease presentation saved to: {PPTX_PATH}")


if __name__ == "__main__":
    create_deck()
